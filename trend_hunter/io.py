from __future__ import annotations

import csv
import json
import logging
from dataclasses import replace
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

from .models import CalendarEvent, Trend, Voice

LOGGER = logging.getLogger(__name__)
REQUIRED_FIELDS = {
    "id",
    "parent_id",
    "platform",
    "channel",
    "user_name",
    "title",
    "content",
    "created_at",
}


def parse_as_of(value: str, timezone: str = "Asia/Hong_Kong") -> datetime:
    """Interpret an as-of date as end-of-day in the challenge timezone."""
    parsed = date.fromisoformat(value)
    return datetime.combine(parsed, time.max, tzinfo=ZoneInfo(timezone))


def _parse_datetime(value: str, timezone: str) -> datetime:
    dt = datetime.fromisoformat(value)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZoneInfo(timezone))
    return dt.astimezone(ZoneInfo(timezone))


def load_voices(path: str | Path, cutoff: datetime, timezone: str) -> list[Voice]:
    """Load only eligible rows. Future rows never leave this function."""
    voices: list[Voice] = []
    seen_ids: set[str] = set()
    with Path(path).open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                LOGGER.warning("Skipping malformed JSON on line %s: %s", line_number, exc)
                continue
            missing = REQUIRED_FIELDS.difference(row)
            if missing:
                LOGGER.warning("Skipping line %s with missing fields: %s", line_number, missing)
                continue
            try:
                created_at = _parse_datetime(str(row["created_at"]), timezone)
            except (TypeError, ValueError) as exc:
                LOGGER.warning("Skipping line %s with bad created_at: %s", line_number, exc)
                continue
            if created_at > cutoff:
                continue
            voice_id = str(row["id"])
            if voice_id in seen_ids:
                LOGGER.warning("Skipping duplicate id %s", voice_id)
                continue
            seen_ids.add(voice_id)
            voices.append(
                Voice(
                    id=voice_id,
                    parent_id=str(row["parent_id"]) if row["parent_id"] else None,
                    platform=str(row["platform"]),
                    channel=str(row["channel"]),
                    user_name=str(row["user_name"]),
                    title=str(row["title"]) if row["title"] else None,
                    content=str(row["content"] or ""),
                    created_at=created_at,
                )
            )
    voices.sort(key=lambda item: (item.created_at, item.id))
    return voices


def _date_value(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip()[:10])


def load_calendar(path: str | Path | None, cutoff: datetime) -> list[CalendarEvent]:
    if path is None:
        return []
    source = Path(path)
    if not source.exists():
        LOGGER.warning("Calendar file not found: %s", source)
        return []
    rows: list[dict[str, object]] = []
    if source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            rows.extend(csv.DictReader(handle))
    elif source.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency error is explicit
            raise RuntimeError("openpyxl is required to read an XLSX calendar") from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip().lower() for value in next(values)]
        rows.extend(dict(zip(headers, row, strict=False)) for row in values)
    else:
        raise ValueError(f"Unsupported calendar format: {source.suffix}")

    events: list[CalendarEvent] = []
    for row in rows:
        if not row.get("date"):
            continue
        event_date = _date_value(row["date"])
        if event_date > cutoff.date():
            continue
        events.append(
            CalendarEvent(
                date=event_date,
                version=str(row.get("version") or "-").strip(),
                notes=str(row.get("notes") or "").strip(),
            )
        )
    return sorted(events, key=lambda event: event.date)


def discover_calendar(input_path: str | Path) -> Path | None:
    directory = Path(input_path).resolve().parent
    for name in ("versions.csv", "versions.xlsx"):
        candidate = directory / name
        if candidate.exists():
            return candidate
    return None


def write_output(
    path: str | Path,
    as_of: str,
    trends: list[Trend],
) -> None:
    payload = {
        "as_of": as_of,
        "trends": [trend.to_dict() for trend in trends],
    }
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    temporary.replace(destination)


def write_audit(path: str | Path, as_of: str, diagnostics: dict[str, object]) -> None:
    payload = {
        "as_of": as_of,
        "stage_counts": diagnostics.get("stage_counts", {}),
        "selected_slots": diagnostics.get("selected_slots", []),
        "candidates": diagnostics.get("selection_audit", []),
    }
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    temporary.replace(destination)


def replace_voice(voice: Voice, **changes: object) -> Voice:
    return replace(voice, **changes)
